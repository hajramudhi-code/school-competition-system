from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Protection
from openpyxl.utils.exceptions import InvalidFileException

from api.exceptions import unprocessable, validation_error
from api.models import Question, Subject
from api.services import validate_question_payload

HEADERS = {
    Question.Mode.MULTIPLE_CHOICE: [
        "text",
        "optionA",
        "optionB",
        "optionC",
        "optionD",
        "correctAnswer",
        "marks",
    ],
    Question.Mode.TRUE_FALSE: ["text", "correctAnswer", "marks"],
    Question.Mode.MENTION: ["text", "correctAnswer", "marks"],
}


def build_template(subject, mode, question_count, marks_per_question):
    headers = HEADERS.get(mode)
    if not headers:
        raise validation_error("Invalid mode.", fields={"mode": "Invalid mode."})
    if question_count < 1:
        raise validation_error("questionCount must be greater than 0.", fields={"questionCount": "Must be > 0."})
    if marks_per_question < 1:
        raise validation_error("marksPerQuestion must be greater than 0.", fields={"marksPerQuestion": "Must be > 0."})

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet["A1"] = f"Subject: {subject.name}"
    sheet["A2"] = f"Mode: {mode}"
    bold = Font(bold=True)
    sheet["A1"].font = bold
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=col, value=header)
        cell.font = bold
    for row in range(5, 5 + question_count):
        sheet.cell(row=row, column=headers.index("marks") + 1, value=marks_per_question)
        if mode == Question.Mode.TRUE_FALSE:
            sheet.cell(row=row, column=headers.index("correctAnswer") + 1, value="TRUE")
        elif mode == Question.Mode.MULTIPLE_CHOICE:
            sheet.cell(row=row, column=headers.index("correctAnswer") + 1, value="A")
        sheet.cell(row=row, column=1, value=f"Question {row - 4}")
        if mode == Question.Mode.MULTIPLE_CHOICE:
            sheet.cell(row=row, column=2, value="Option A")
            sheet.cell(row=row, column=3, value="Option B")
            sheet.cell(row=row, column=4, value="Option C")
            sheet.cell(row=row, column=5, value="Option D")
        elif mode == Question.Mode.MENTION:
            sheet.cell(row=row, column=2, value="Expected answer")

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    safe_name = subject.name.lower().replace(" ", "_")
    filename = f"template_{safe_name}_{mode.lower()}.xlsx"
    return buffer, filename


def parse_upload(file, subject, mode):
    expected = HEADERS.get(mode)
    if not expected:
        raise validation_error("Invalid mode.", fields={"mode": "Invalid mode."})
    try:
        workbook = load_workbook(filename=file, data_only=True)
    except (InvalidFileException, OSError, ValueError) as exc:
        raise validation_error("Unreadable or malformed file.") from exc

    sheet = workbook.active
    header_row = None
    for row in sheet.iter_rows(min_row=1, max_row=10, values_only=False):
        values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
        if values and values[0] == expected[0]:
            header_row = [cell.value for cell in row if cell.value]
            break
    if not header_row:
        raise unprocessable("Schema mismatch with selected mode.", fields={"headers": expected})

    normalized = [str(h).strip() for h in header_row]
    if normalized[: len(expected)] != expected:
        raise unprocessable("Schema mismatch with selected mode.", fields={"expected": expected, "found": normalized})

    total_rows = 0
    imported = []
    missing_fields = []
    invalid = 0
    duplicates = 0
    seen_text = set(
        Question.objects.filter(subject=subject, mode=mode).values_list("text", flat=True)
    )

    for excel_row in sheet.iter_rows(min_row=5, values_only=True):
        if not excel_row or all(cell is None or str(cell).strip() == "" for cell in excel_row):
            continue
        total_rows += 1
        row_number = total_rows + 4
        mapping = {expected[i]: excel_row[i] if i < len(excel_row) else None for i in range(len(expected))}
        missing = [field for field in expected if mapping.get(field) in (None, "")]
        if missing:
            invalid += 1
            missing_fields.append({"row": row_number, "fields": missing})
            continue
        text = str(mapping["text"]).strip()
        if text in seen_text:
            duplicates += 1
            continue
        payload = {
            "subjectId": str(subject.id),
            "mode": mode,
            "text": text,
            "marks": mapping["marks"],
            "correctAnswer": mapping["correctAnswer"],
            "isVideoQuestion": False,
        }
        if mode == Question.Mode.MULTIPLE_CHOICE:
            payload.update(
                {
                    "optionA": mapping["optionA"],
                    "optionB": mapping["optionB"],
                    "optionC": mapping["optionC"],
                    "optionD": mapping["optionD"],
                }
            )
        try:
            cleaned = validate_question_payload(payload)
        except Exception:
            invalid += 1
            missing_fields.append({"row": row_number, "fields": ["invalid"]})
            continue
        question = Question.objects.create(
            subject=cleaned["subject"],
            mode=cleaned["mode"],
            text=cleaned["text"],
            option_a=cleaned["option_a"],
            option_b=cleaned["option_b"],
            option_c=cleaned["option_c"],
            option_d=cleaned["option_d"],
            correct_answer=cleaned["correct_answer"],
            marks=cleaned["marks"],
            is_video_question=False,
        )
        imported.append(str(question.id))
        seen_text.add(text)

    return {
        "totalRows": total_rows,
        "validQuestions": len(imported),
        "invalidQuestions": invalid,
        "duplicateQuestions": duplicates,
        "missingFields": missing_fields,
        "importedQuestionIds": imported,
    }
